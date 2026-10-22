import csv
import json
from datetime import datetime
from pathlib import Path

from core.config.settings import config
from core.models.salary import Salary
from core.repositories.salary_repository import SalaryRepository
from core.services.export_service import ExportService
from core.services.extra_payment_service import ExtraPaymentService


def _seed_salary_row():
    SalaryRepository.save(
        Salary(
            year=2025,
            month=9,
            advance=11111.0,
            salary=22222.0,
            advance_date=datetime.strptime("10.09.2025", "%d.%m.%Y"),
            salary_date=datetime.strptime("25.09.2025", "%d.%m.%Y"),
        )
    )


def test_export_to_csv_creates_file_with_data(temp_db, tmp_path: Path):
    old_export_dir = config.export.export_dir
    config.export.export_dir = str(tmp_path)
    try:
        _seed_salary_row()
        filepath = ExportService.export_to_csv()
        file_path_obj = Path(filepath)

        assert file_path_obj.exists()
        assert file_path_obj.suffix == ".csv"
        # BOM: по нему Excel узнаёт UTF-8 и не показывает кириллицу «кракозябрами»
        assert file_path_obj.read_bytes().startswith(b"\xef\xbb\xbf")

        with file_path_obj.open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f, delimiter=";"))

        assert len(rows) >= 2
        assert rows[0][0] == "Год"
        assert rows[1][0] == "2025"
        assert rows[1][1] == "Сентябрь"
        # Суммы с запятой — русский Excel читает их как числа
        assert rows[0][6:8] == ["Доп. выплаты", "Итого"]
        assert (rows[1][2], rows[1][4], rows[1][6], rows[1][7]) == ("11111,00", "22222,00", "0,00", "33333,00")
    finally:
        config.export.export_dir = old_export_dir


def test_export_to_json_creates_structured_payload(temp_db, tmp_path: Path):
    old_export_dir = config.export.export_dir
    config.export.export_dir = str(tmp_path)
    try:
        _seed_salary_row()
        filepath = ExportService.export_to_json()
        file_path_obj = Path(filepath)

        assert file_path_obj.exists()
        assert file_path_obj.suffix == ".json"

        payload = json.loads(file_path_obj.read_text(encoding="utf-8"))
        assert "export_date" in payload
        assert payload["total_records"] >= 1
        assert isinstance(payload["data"], list)
        assert payload["data"][0]["year"] == 2025
    finally:
        config.export.export_dir = old_export_dir


def _export_excel(tmp_path: Path):
    from openpyxl import load_workbook

    old_export_dir = config.export.export_dir
    config.export.export_dir = str(tmp_path)
    try:
        filepath = Path(ExportService.export_to_excel())
    finally:
        config.export.export_dir = old_export_dir
    assert filepath.suffix == ".xlsx"
    return load_workbook(filepath)


def test_export_to_excel_writes_sheets_with_data(temp_db, tmp_path: Path):
    _seed_salary_row()
    ExtraPaymentService.add(2025, 9, "bonus", 5000.0, "15.09.2025", "за квартал")

    workbook = _export_excel(tmp_path)

    assert workbook.sheetnames == ["Данные", "Сводка по годам", "Статистика", "Доп. выплаты"]
    data = workbook["Данные"]
    assert [cell.value for cell in data[1]] == [
        "Год", "Месяц", "Аванс", "Дата аванса", "Зарплата",
        "Дата зарплаты", "Доп. выплаты", "Итого", "Дата создания (МСК)", "Заметка",
    ]
    assert [cell.value for cell in data[2]][:8] == [
        2025, 9, 11111.0, "2025-09-10", 22222.0, "2025-09-25", 5000.0, 38333.0,
    ]
    yearly = [cell.value for cell in workbook["Сводка по годам"][2]]
    assert yearly == [2025, 11111.0, 22222.0, 5000.0, 38333.0, 1]
    stats = [cell.value for cell in workbook["Статистика"][2]]
    assert stats == [1, 11111.0, 22222.0, 5000.0, 38333.0, 2025, 2025]
    extras = [cell.value for cell in workbook["Доп. выплаты"][2]]
    assert extras == [2025, 9, "Премия", 5000.0, "2025-09-15", "за квартал"]


def test_export_to_excel_keeps_header_style(temp_db, tmp_path: Path):
    _seed_salary_row()

    header = _export_excel(tmp_path)["Данные"]["A1"]

    assert header.font.b is True
    assert header.border.left.style == "thin" and header.border.bottom.style == "thin"
    assert (header.alignment.horizontal, header.alignment.vertical) == ("center", "top")


def test_export_to_excel_works_on_empty_database(temp_db, tmp_path: Path):
    # Раньше суммы NULL из пустой таблицы складывались и экспорт падал с TypeError.
    workbook = _export_excel(tmp_path)

    assert workbook["Данные"].max_row == 1  # только заголовок
    assert [cell.value for cell in workbook["Статистика"][2]] == [0, 0, 0, 0, 0, None, None]
    assert workbook["Доп. выплаты"].max_row == 1


def _set_created_at(utc_text: str) -> None:
    from core.database.connection import DatabaseConnection

    conn = DatabaseConnection.get_connection()
    conn.execute("UPDATE salaries SET created_at = ?", (utc_text,))
    conn.commit()


def test_export_shows_created_at_in_moscow(temp_db, tmp_path: Path):
    # created_at пишет SQLite в UTC; 21:38 UTC — это уже следующий день по МСК
    _seed_salary_row()
    _set_created_at("2026-10-14 21:38:43")
    old_export_dir = config.export.export_dir
    config.export.export_dir = str(tmp_path)
    try:
        csv_path = Path(ExportService.export_to_csv())
    finally:
        config.export.export_dir = old_export_dir

    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert rows[0][8] == "Дата создания (МСК)"
    assert rows[1][8] == "2026-10-15 00:38:43"
    assert _export_excel(tmp_path)["Данные"]["I2"].value == "2026-10-15 00:38:43"


def test_json_export_date_is_utc(temp_db, tmp_path: Path):
    old_export_dir = config.export.export_dir
    config.export.export_dir = str(tmp_path)
    try:
        payload = json.loads(Path(ExportService.export_to_json()).read_text(encoding="utf-8"))
    finally:
        config.export.export_dir = old_export_dir

    assert payload["export_date"].endswith("+00:00")
